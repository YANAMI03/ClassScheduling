import io
import pytest
import openpyxl
import app as app_module
import excel_export


def test_parse_time_preserves_pm_and_am():
    """Verify that PM and AM times are parsed correctly with zero AM/PM confusion."""
    # 1:00 PM -> 13:00 (46800 sec)
    sec_1pm = excel_export._parse_time_to_seconds('01:00 PM')
    assert sec_1pm == 13 * 3600

    # 1:00 AM -> 01:00 (3600 sec)
    sec_1am = excel_export._parse_time_to_seconds('01:00 AM')
    assert sec_1am == 1 * 3600
    assert sec_1pm != sec_1am

    # 12:00 PM -> 12:00 (43200 sec)
    sec_12pm = excel_export._parse_time_to_seconds('12:00 PM')
    assert sec_12pm == 12 * 3600

    # 12:00 AM -> 00:00 (0 sec)
    sec_12am = excel_export._parse_time_to_seconds('12:00 AM')
    assert sec_12am == 0

    # Test reverse display formatting
    assert excel_export._seconds_to_display_time(sec_1pm) == '01:00 PM'
    assert excel_export._seconds_to_display_time(sec_1am) == '01:00 AM'
    assert excel_export._seconds_to_display_time(sec_12pm) == '12:00 PM'


def test_generate_excel_room_schedule():
    """Verify room timetable Excel generation with vertical stretching and formatting."""
    room = {'room_id': 1, 'room_name': 'Room 101', 'room_type': 'Lecture'}
    entries = [
        {
            'day': 'Monday',
            'start_time_raw': '08:00:00',
            'end_time_raw': '10:00:00',
            'course_code': 'IT101 - Intro to Computing',
            'section': '1A',
            'professor': 'Dr. Alan Turing',
            'session_type': 'Lecture',
        }
    ]
    buf = excel_export.generate_timetable_excel('room', room, entries, filter_metadata={'semester': '1st Semester'})
    wb = openpyxl.load_workbook(buf)
    ws = wb.active

    # Check Title and Subtitle
    assert ws['A2'].value == 'ROOM SCHEDULE'
    assert 'Room 101' in ws['A3'].value

    # Check Headers
    headers = [ws.cell(row=6, column=c).value for c in range(1, 8)]
    assert headers == ['TIME', 'MONDAY', 'TUESDAY', 'WEDNESDAY', 'THURSDAY', 'FRIDAY', 'SATURDAY']

    # Check Vertical Stretching (row heights >= 65 pt)
    for r in range(7, ws.max_row + 1):
        assert ws.row_dimensions[r].height >= 65

    # Check Column widths
    assert ws.column_dimensions['A'].width >= 18
    assert ws.column_dimensions['B'].width >= 24

    # Check Print Orientation
    assert ws.page_setup.orientation == 'landscape'


def test_generate_excel_section_schedule():
    """Verify section timetable Excel generation, multi-hour merging, and content."""
    section = {'section_name': 'BSIT 2A', 'year_level': '2', 'semester': '2nd Semester', 'major': 'Network'}
    entries = [
        {
            'day': 'Tuesday',
            'start_time_raw': '09:00:00',
            'end_time_raw': '11:00:00',
            'course_code': 'NET201 - Advanced Networking',
            'professor': 'Grace Hopper',
            'room': 'Cisco Lab',
            'session_type': 'Laboratory',
        }
    ]
    buf = excel_export.generate_timetable_excel('section', section, entries, filter_metadata={'semester': '2nd Semester', 'year': '2', 'major': 'Network'})
    wb = openpyxl.load_workbook(buf)
    ws = wb.active

    assert ws['A2'].value == 'SECTION SCHEDULE'
    assert 'BSIT 2A' in ws['A3'].value

    # Tuesday is column C (Col 3)
    # Find the cell containing the class
    found = False
    for r in range(7, ws.max_row + 1):
        cell_val = str(ws.cell(row=r, column=3).value or '')
        if 'NET201' in cell_val:
            found = True
            assert 'Grace Hopper' in cell_val
            assert 'Cisco Lab' in cell_val
            assert ws.cell(row=r, column=3).alignment.wrap_text is True
            assert ws.cell(row=r, column=3).alignment.vertical == 'center'
            break
    assert found, "NET201 class not found in Tuesday column"


def test_generate_excel_professor_schedule():
    """Verify professor timetable Excel generation."""
    professor = {'prof_id': 1, 'first_name': 'Ada', 'last_name': 'Lovelace', 'department': 'Computer Science'}
    entries = [
        {
            'day': 'Wednesday',
            'start_time_raw': '13:00:00',
            'end_time_raw': '15:00:00',
            'course_code': 'CS102 - Data Structures',
            'section': '1B',
            'room': 'Room 305',
            'session_type': 'Lecture',
        }
    ]
    buf = excel_export.generate_timetable_excel('professor', professor, entries, filter_metadata={'semester': '1st Semester'})
    wb = openpyxl.load_workbook(buf)
    ws = wb.active

    assert ws['A2'].value == 'FACULTY PROGRAM'
    assert 'Ada Lovelace' in ws['A3'].value

    # Wednesday is column D (Col 4)
    found = False
    for r in range(7, ws.max_row + 1):
        cell_val = str(ws.cell(row=r, column=4).value or '')
        if 'CS102' in cell_val:
            found = True
            assert 'Section: BSIT - 1B' in cell_val
            assert 'Room 305' in cell_val
            assert '01:00 PM - 03:00 PM' in cell_val
            break
    assert found, "CS102 class not found in Wednesday column"


def test_generate_excel_multiple_classes_same_slot():
    """Requirement: If multiple classes occur on the same day/time, display them properly without losing information."""
    section = {'section_name': 'Irregular-Combo'}
    entries = [
        {
            'day': 'Monday',
            'start_time_raw': '08:00:00',
            'end_time_raw': '09:00:00',
            'course_code': 'MATH101 - Calculus 1',
            'professor': 'Prof Gauss',
            'room': 'Room 101',
            'session_type': 'Lecture',
        },
        {
            'day': 'Monday',
            'start_time_raw': '08:00:00',
            'end_time_raw': '09:00:00',
            'course_code': 'HIST101 - Philippine History',
            'professor': 'Prof Rizal',
            'room': 'Room 102',
            'session_type': 'Lecture',
        }
    ]
    buf = excel_export.generate_timetable_excel('section', section, entries)
    wb = openpyxl.load_workbook(buf)
    ws = wb.active

    # Monday is column B (Col 2)
    found_both = False
    for r in range(7, ws.max_row + 1):
        cell_val = str(ws.cell(row=r, column=2).value or '')
        if 'MATH101' in cell_val and 'HIST101' in cell_val:
            found_both = True
            assert 'Prof Gauss' in cell_val
            assert 'Prof Rizal' in cell_val
            break
    assert found_both, "Both classes should appear in the cell without losing information"


def test_routes_export_room_excel(monkeypatch):
    """Test GET /room_schedule/<room_id>/export returns valid Excel file."""
    client = app_module.app.test_client()

    with client.session_transaction() as sess:
        sess['user_id'] = 'test-user-id'
        sess['role'] = 'admin'
        sess['email'] = 'admin@example.com'

    class FakeQuery:
        def __init__(self, table_name):
            self.table_name = table_name

        def select(self, *args, **kwargs):
            return self

        def eq(self, *args, **kwargs):
            return self

        def like(self, *args, **kwargs):
            return self

        def execute(self):
            if self.table_name == 'room':
                return type('Resp', (), {'data': [{'room_id': 10, 'room_name': 'CL-1', 'room_type': 'Laboratory'}]})()
            elif self.table_name == 'schedule':
                return type('Resp', (), {'data': [
                    {
                        'schedule_id': 1,
                        'professor_load_id': 1,
                        'room_id': 10,
                        'day': 'Monday',
                        'class_start': '08:00:00',
                        'class_end': '10:00:00',
                        'section': '1A',
                        'semester': '1st Semester',
                        'major': None,
                        'session_type': 'Laboratory',
                        'professor_load': {
                            'course': {'course_id': 1, 'course_code': 'IT101 - Intro to Computing'},
                            'professor': {'prof_id': 1, 'first_name': 'Alan', 'last_name': 'Turing'}
                        },
                        'room': {'room_name': 'CL-1'}
                    }
                ]})()
            elif self.table_name == 'working_hours':
                return type('Resp', (), {'data': [{'start_day': 'Monday', 'end_day': 'Saturday', 'start_time': '07:00:00', 'end_time': '19:00:00', 'lunch_time': '12:00:00'}]})()
            return type('Resp', (), {'data': []})()

    monkeypatch.setattr(app_module.supabase, 'table', lambda name: FakeQuery(name))

    resp = client.get('/room_schedule/10/export?semester=1st+Semester')
    assert resp.status_code == 200
    assert 'application/vnd.openxmlformats-officedocument.spreadsheetml.sheet' in resp.content_type
    assert 'attachment;' in resp.headers.get('Content-Disposition', '')
    assert 'Room_Schedule_CL-1.xlsx' in resp.headers.get('Content-Disposition', '')

    wb = openpyxl.load_workbook(io.BytesIO(resp.data))
    assert wb.active['A2'].value == 'ROOM SCHEDULE'


def test_routes_export_section_excel(monkeypatch):
    """Test GET /schedule/<section_name>/export returns valid Excel file."""
    client = app_module.app.test_client()

    with client.session_transaction() as sess:
        sess['user_id'] = 'test-user-id'
        sess['role'] = 'admin'
        sess['email'] = 'admin@example.com'

    class FakeQuery:
        def __init__(self, table_name):
            self.table_name = table_name

        def select(self, *args, **kwargs):
            return self

        def eq(self, *args, **kwargs):
            return self

        def or_(self, *args, **kwargs):
            return self

        def execute(self):
            if self.table_name == 'schedule':
                return type('Resp', (), {'data': [
                    {
                        'schedule_id': 2,
                        'day': 'Wednesday',
                        'class_start': '10:00:00',
                        'class_end': '12:00:00',
                        'session_type': 'Lecture',
                        'semester': '1st Semester',
                        'major': 'General',
                        'professor_load': {
                            'course': {'course_id': 2, 'course_code': 'CC-102 Programming'},
                            'professor': {'prof_id': 2, 'first_name': 'Grace', 'last_name': 'Hopper'}
                        },
                        'room': {'room_name': 'Room 201'}
                    }
                ]})()
            elif self.table_name == 'working_hours':
                return type('Resp', (), {'data': []})()
            return type('Resp', (), {'data': []})()

    monkeypatch.setattr(app_module.supabase, 'table', lambda name: FakeQuery(name))

    resp = client.get('/schedule/1A/export?semester=1st+Semester')
    assert resp.status_code == 200
    assert 'application/vnd.openxmlformats-officedocument.spreadsheetml.sheet' in resp.content_type
    assert 'attachment;' in resp.headers.get('Content-Disposition', '')
    assert 'Section_Schedule_1A.xlsx' in resp.headers.get('Content-Disposition', '')

    wb = openpyxl.load_workbook(io.BytesIO(resp.data))
    assert wb.active['A2'].value == 'SECTION SCHEDULE'


def test_routes_export_professor_excel(monkeypatch):
    """Test GET /professor_schedule/<professor_id>/export returns valid Excel file."""
    client = app_module.app.test_client()

    with client.session_transaction() as sess:
        sess['user_id'] = 'test-user-id'
        sess['role'] = 'admin'
        sess['email'] = 'admin@example.com'

    class FakeQuery:
        def __init__(self, table_name):
            self.table_name = table_name

        def select(self, *args, **kwargs):
            return self

        def eq(self, *args, **kwargs):
            return self

        def like(self, *args, **kwargs):
            return self

        def execute(self):
            if self.table_name == 'professor':
                return type('Resp', (), {'data': [{'prof_id': 5, 'first_name': 'Alan', 'last_name': 'Turing', 'department': 'CS', 'max_hours': 30}]})()
            elif self.table_name == 'professor_load':
                return type('Resp', (), {'data': [{'professor_load_id': 10, 'course_id': 1, 'course': {'course_code': 'IT101'}}]})()
            elif self.table_name == 'schedule':
                return type('Resp', (), {'data': [
                    {
                        'schedule_id': 3,
                        'professor_load_id': 10,
                        'room_id': 1,
                        'day': 'Friday',
                        'class_start': '14:00:00',
                        'class_end': '17:00:00',
                        'section': '2B',
                        'semester': '1st Semester',
                        'major': None,
                        'session_type': 'Lecture',
                        'professor_load': {'course': {'course_code': 'IT101'}},
                        'room': {'room_name': 'Room 105'}
                    }
                ]})()
            elif self.table_name == 'working_hours':
                return type('Resp', (), {'data': []})()
            return type('Resp', (), {'data': []})()

    monkeypatch.setattr(app_module.supabase, 'table', lambda name: FakeQuery(name))

    resp = client.get('/professor_schedule/5/export')
    assert resp.status_code == 200
    assert 'application/vnd.openxmlformats-officedocument.spreadsheetml.sheet' in resp.content_type
    assert 'attachment;' in resp.headers.get('Content-Disposition', '')
    assert 'Faculty_Program_Alan_Turing.xlsx' in resp.headers.get('Content-Disposition', '')

    wb = openpyxl.load_workbook(io.BytesIO(resp.data))
    assert wb.active['A2'].value == 'FACULTY PROGRAM'


def test_ui_templates_contain_export_to_excel_buttons(monkeypatch):
    """Test that the schedule pages render with the 'Export to Excel' buttons."""
    client = app_module.app.test_client()

    with client.session_transaction() as sess:
        sess['user_id'] = 'test-user-id'
        sess['role'] = 'admin'
        sess['email'] = 'admin@example.com'

    class FakeQuery:
        def __init__(self, table_name):
            self.table_name = table_name

        def select(self, *args, **kwargs):
            return self

        def eq(self, *args, **kwargs):
            return self

        def or_(self, *args, **kwargs):
            return self

        def like(self, *args, **kwargs):
            return self

        def execute(self):
            if self.table_name == 'room':
                return type('Resp', (), {'data': [{'room_id': 1, 'room_name': '101', 'room_type': 'Lecture'}]})()
            elif self.table_name == 'professor':
                return type('Resp', (), {'data': [{'prof_id': 1, 'first_name': 'Alan', 'last_name': 'Turing', 'department': 'CS', 'max_hours': 30}]})()
            elif self.table_name == 'professor_load':
                return type('Resp', (), {'data': []})()
            elif self.table_name == 'working_hours':
                return type('Resp', (), {'data': []})()
            elif self.table_name == 'schedule':
                return type('Resp', (), {'data': []})()
            return type('Resp', (), {'data': []})()

    monkeypatch.setattr(app_module.supabase, 'table', lambda name: FakeQuery(name))

    # Room schedule details page
    resp_room = client.get('/room_schedule/1')
    assert resp_room.status_code == 200
    assert b'Export to Excel' in resp_room.data
    assert b'/room_schedule/1/export' in resp_room.data

    # Section schedule details page
    resp_sec = client.get('/schedule/1A')
    assert resp_sec.status_code == 200
    assert b'Export to Excel' in resp_sec.data
    assert b'/schedule/1A/export' in resp_sec.data

    # Professor schedule details page
    resp_prof = client.get('/professor_schedule/1')
    assert resp_prof.status_code == 200
    assert b'Export to Excel' in resp_prof.data
    assert b'/professor_schedule/1/export' in resp_prof.data


def test_all_7_themes_generation_and_styling():
    """Verify that each of the 7 themes applies its distinctive header background, white text,
    proper borders, and sheet tab color without breaking the timetable structure.
    """
    expected_themes = ['Blue', 'Red', 'Green', 'Purple', 'Orange', 'Teal', 'Pink']
    assert set(expected_themes) == set(excel_export.EXCEL_THEMES.keys())

    entries = [
        {
            'day': 'Monday',
            'start_time_raw': '08:00:00',
            'end_time_raw': '10:00:00',
            'course_code': 'CS101 - Intro to CS',
            'section': '1A',
            'professor': 'Prof Turing',
            'room': 'Lab 1',
            'session_type': 'Laboratory',
        }
    ]

    for theme_name in expected_themes:
        palette = excel_export.EXCEL_THEMES[theme_name]
        buf = excel_export.generate_timetable_excel(
            'section',
            {'section_name': f'Section {theme_name}'},
            entries,
            theme=theme_name
        )
        wb = openpyxl.load_workbook(buf)
        ws = wb.active

        # 1. Structure must be intact
        headers = [ws.cell(row=6, column=c).value for c in range(1, 8)]
        assert headers == ['TIME', 'MONDAY', 'TUESDAY', 'WEDNESDAY', 'THURSDAY', 'FRIDAY', 'SATURDAY']

        # 2. Main headers (Row 6) must use theme header_bg and header_fg (pure white)
        for c in range(1, 8):
            cell = ws.cell(row=6, column=c)
            # openpyxl fills might include alpha prefix '00' or direct hex
            cell_fill_rgb = str(cell.fill.start_color.rgb or '').upper()
            assert palette['header_bg'].upper() in cell_fill_rgb, (
                f"Theme {theme_name} header cell ({c}) should have bg {palette['header_bg']}, got {cell_fill_rgb}"
            )
            cell_font_rgb = str(cell.font.color.rgb or '').upper()
            assert palette['header_fg'].upper() in cell_font_rgb, (
                f"Theme {theme_name} header cell ({c}) should have fg {palette['header_fg']}, got {cell_font_rgb}"
            )

        # 3. Document title should carry theme header_bg
        title_font_rgb = str(ws['A2'].font.color.rgb or '').upper()
        assert palette['header_bg'].upper() in title_font_rgb

        # 4. Sheet Tab Color must match theme header_bg
        tab_rgb = str(ws.sheet_properties.tabColor.rgb or '').upper()
        assert palette['header_bg'].upper() in tab_rgb

        # 5. Class cell fill must match theme class_fill
        found_class = False
        for r in range(7, ws.max_row + 1):
            if 'CS101' in str(ws.cell(row=r, column=2).value or ''):
                found_class = True
                class_fill_rgb = str(ws.cell(row=r, column=2).fill.start_color.rgb or '').upper()
                assert palette['class_fill'].upper() in class_fill_rgb
                break
        assert found_class, f"CS101 class not found in Monday column for theme {theme_name}"

        # 6. Row height vertical stretching intact
        for r in range(7, ws.max_row + 1):
            assert ws.row_dimensions[r].height >= 65


def test_ui_theme_dropdown_removed(monkeypatch):
    """Verify UI templates do NOT render the 'Export Theme:' dropdown."""
    client = app_module.app.test_client()
    with client.session_transaction() as sess:
        sess['user_id'] = 'test-admin'
        sess['role'] = 'admin'

    class FakeQuery:
        def __init__(self, table_name):
            self.table_name = table_name
        def select(self, *args, **kwargs): return self
        def eq(self, *args, **kwargs): return self
        def or_(self, *args, **kwargs): return self
        def like(self, *args, **kwargs): return self
        def execute(self):
            return type('Resp', (), {'data': []})()

    monkeypatch.setattr(app_module.supabase, 'table', lambda name: FakeQuery(name))

    resp_sec = client.get('/schedule/1A')
    assert resp_sec.status_code == 200
    assert b'Export Theme:' not in resp_sec.data
    assert b'exportThemeSelect' not in resp_sec.data
    assert b'Export to Excel' in resp_sec.data


def test_export_section_schedule_clean_export(monkeypatch):
    """Verify GET /schedule/<section>/export exports without theme selector."""
    client = app_module.app.test_client()
    with client.session_transaction() as sess:
        sess['user_id'] = 'test-admin'
        sess['role'] = 'admin'

    class FakeQuery:
        def __init__(self, table_name):
            self.table_name = table_name
        def select(self, *args, **kwargs): return self
        def eq(self, *args, **kwargs): return self
        def or_(self, *args, **kwargs): return self
        def like(self, *args, **kwargs): return self
        def execute(self):
            if self.table_name == 'schedule':
                return type('Resp', (), {'data': [
                    {
                        'schedule_id': 1,
                        'day': 'Monday',
                        'class_start': '08:00:00',
                        'class_end': '10:00:00',
                        'session_type': 'Lecture',
                        'semester': '1st Semester',
                        'major': None,
                        'professor_load': {
                            'course': {'course_code': 'CS101 Intro'},
                            'professor': {'first_name': 'Alan', 'last_name': 'Turing'}
                        },
                        'room': {'room_name': 'R101'}
                    }
                ]})()
            return type('Resp', (), {'data': []})()

    monkeypatch.setattr(app_module.supabase, 'table', lambda name: FakeQuery(name))

    resp = client.get('/schedule/1A/export')
    assert resp.status_code == 200
    wb = openpyxl.load_workbook(io.BytesIO(resp.data))
    ws = wb.active
    assert ws['A2'].value == 'SECTION SCHEDULE'
    assert '1A' in ws['A3'].value


def test_excel_export_faculty_program_title_and_section_format():
    """Verify that excel export uses 'FACULTY PROGRAM' heading and formats sections with program."""
    assert excel_export._format_section_with_program('4A', 'BSIT') == 'BSIT - 4A'
    assert excel_export._format_section_with_program('4A', None) == 'BSIT - 4A'
    assert excel_export._format_section_with_program('BSIT - 4A', 'BSIT') == 'BSIT - 4A'
    assert excel_export._format_section_with_program('BSIT 4A', 'BSIT') == 'BSIT 4A'
    assert excel_export._format_section_with_program('3B', 'BSDS') == 'BSDS - 3B'

    # Professor schedule excel generation heading
    prof = {'prof_id': 10, 'first_name': 'Alan', 'last_name': 'Turing'}
    entries = [
        {'day': 'Monday', 'start_time_raw': '08:00:00', 'end_time_raw': '10:00:00', 'course_code': 'IT101', 'section': '4A', 'program': 'BSIT', 'room': 'Lab 1', 'session_type': 'Lecture'}
    ]
    buf = excel_export.generate_timetable_excel('professor', prof, entries)
    wb = openpyxl.load_workbook(buf)
    ws = wb.active
    assert ws['A2'].value == 'FACULTY PROGRAM'




