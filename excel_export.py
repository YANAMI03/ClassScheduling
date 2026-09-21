import io
import os
import json
import re
from datetime import datetime, timedelta
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter


# ==============================================================================
# 7 PREDEFINED COLOR THEMES FOR EXCEL TIMETABLE EXPORT
# ==============================================================================

EXCEL_THEMES = {
    'Blue': {
        'name': 'Blue',
        'display_name': 'Blue Theme',
        'header_bg': '1E3A8A',       # Professional Royal Navy Blue
        'header_fg': 'FFFFFF',       # Crisp White
        'accent_fg': '1D4ED8',       # Dark Royal Blue for subtitles/accents
        'class_fill': 'EFF6FF',      # Light Ice Blue
        'meta_bar_fill': 'DBEAFE',   # Soft light blue tint
        'border_color': 'BFDBFE',    # Soft blue border
        'header_border': '172554',   # Strong header border
    },
    'Red': {
        'name': 'Red',
        'display_name': 'Red Theme',
        'header_bg': '991B1B',       # Professional Crimson / Dark Red
        'header_fg': 'FFFFFF',
        'accent_fg': 'B91C1C',       # Deep Red for subtitles/accents
        'class_fill': 'FEF2F2',      # Soft Blush
        'meta_bar_fill': 'FEE2E2',   # Soft light rose tint
        'border_color': 'FECACA',    # Soft red border
        'header_border': '7F1D1D',
    },
    'Green': {
        'name': 'Green',
        'display_name': 'Green Theme',
        'header_bg': '166534',       # Professional Forest / Emerald Green
        'header_fg': 'FFFFFF',
        'accent_fg': '15803D',       # Deep Green for subtitles/accents
        'class_fill': 'F0FDF4',      # Soft Mint
        'meta_bar_fill': 'DCFCE7',   # Soft light mint tint
        'border_color': 'BBF7D0',    # Soft green border
        'header_border': '14532D',
    },
    'Purple': {
        'name': 'Purple',
        'display_name': 'Purple Theme',
        'header_bg': '581C87',       # Professional Royal Purple / Indigo
        'header_fg': 'FFFFFF',
        'accent_fg': '6B21A8',       # Deep Purple for subtitles/accents
        'class_fill': 'FAF5FF',      # Soft Lavender
        'meta_bar_fill': 'F3E8FF',   # Soft light purple tint
        'border_color': 'E9D5FF',    # Soft purple border
        'header_border': '3B0764',
    },
    'Orange': {
        'name': 'Orange',
        'display_name': 'Orange Theme',
        'header_bg': 'C2410C',       # Professional Rust / Amber Orange
        'header_fg': 'FFFFFF',
        'accent_fg': 'EA580C',       # Deep Amber for subtitles/accents
        'class_fill': 'FFF7ED',      # Soft Peach / Cream
        'meta_bar_fill': 'FFEDD5',   # Soft light orange tint
        'border_color': 'FED7AA',    # Soft orange border
        'header_border': '9A3412',
    },
    'Teal': {
        'name': 'Teal',
        'display_name': 'Teal Theme',
        'header_bg': '115E59',       # Professional Deep Ocean Teal
        'header_fg': 'FFFFFF',
        'accent_fg': '0F766E',       # Deep Teal for subtitles/accents
        'class_fill': 'F0FDFA',      # Soft Seafoam
        'meta_bar_fill': 'CCFBF1',   # Soft light teal tint
        'border_color': '99F6E4',    # Soft teal border
        'header_border': '134E4A',
    },
    'Pink': {
        'name': 'Pink',
        'display_name': 'Pink Theme',
        'header_bg': '9D174D',       # Professional Mulberry / Rose Pink
        'header_fg': 'FFFFFF',
        'accent_fg': 'BE185D',       # Deep Rose for subtitles/accents
        'class_fill': 'FDF2F8',      # Soft Rose Quartz
        'meta_bar_fill': 'FCE7F3',   # Soft light rose tint
        'border_color': 'FBCFE8',    # Soft pink border
        'header_border': '831843',
    },
}

THEME_NAMES = list(EXCEL_THEMES.keys())
DEFAULT_THEME = 'Blue'

# Persistent storage file for section themes
_THEMES_FILE = os.path.join(os.path.dirname(os.path.abspath(__file__)), 'section_export_themes.json')


def _normalize_sec_key(key: str) -> str:
    """Normalize a section key for robust matching (e.g. 'Section 1A' -> '1a')."""
    s = str(key or '').strip().lower()
    if s.startswith('section '):
        s = s[8:].strip()
    elif s.startswith('section'):
        s = s[7:].strip()
    elif s.startswith('sec '):
        s = s[4:].strip()
    return s


def get_all_section_themes():
    """Retrieve all persisted section theme mappings."""
    if not os.path.exists(_THEMES_FILE):
        return {}
    try:
        with open(_THEMES_FILE, 'r', encoding='utf-8') as f:
            data = json.load(f)
            return data if isinstance(data, dict) else {}
    except Exception:
        return {}


def get_section_theme(section_name: str) -> str:
    """Get the assigned Excel export theme for a section, defaulting to 'Blue'.
    Guarantees deterministic, persistent theme assignment without random changes.
    """
    if not section_name:
        return DEFAULT_THEME
    themes = get_all_section_themes()
    sec_key = str(section_name).strip()

    # 1. Exact match
    val = themes.get(sec_key)
    if val in EXCEL_THEMES:
        return val

    # 2. Case-insensitive & normalized match ('Section 1A' vs '1A')
    norm_target = _normalize_sec_key(sec_key)
    for k, v in themes.items():
        if k.strip().lower() == sec_key.lower() and v in EXCEL_THEMES:
            return v
        if _normalize_sec_key(k) == norm_target and v in EXCEL_THEMES:
            return v

    return DEFAULT_THEME


def set_section_theme(section_name: str, theme: str) -> str:
    """Save the assigned Excel export theme for a section to persistent storage."""
    if not section_name:
        return DEFAULT_THEME
    # Normalize theme name to one of the 7 predefined themes
    selected_theme = DEFAULT_THEME
    for t_name in EXCEL_THEMES.keys():
        if t_name.lower() == str(theme).strip().lower():
            selected_theme = t_name
            break

    sec_key = str(section_name).strip()
    themes = get_all_section_themes()
    themes[sec_key] = selected_theme
    try:
        with open(_THEMES_FILE, 'w', encoding='utf-8') as f:
            json.dump(themes, f, indent=2)
    except Exception as e:
        print(f"Warning: Could not save section theme: {e}")
    return selected_theme


def _parse_time_to_seconds(val):
    """Convert any time-like value to total seconds from midnight.
    Correctly handles timedelta, time, 12-hour strings ('01:00 PM'), and 24-hour strings ('13:00:00').
    """
    if val is None:
        return None
    if isinstance(val, timedelta):
        return int(val.total_seconds())
    if hasattr(val, 'hour'):
        return val.hour * 3600 + val.minute * 60 + getattr(val, 'second', 0)

    val_str = str(val).strip().upper()
    if not val_str:
        return None

    is_pm = 'PM' in val_str
    is_am = 'AM' in val_str

    clean = val_str.replace('AM', '').replace('PM', '').strip()
    if ' ' in clean:
        clean = clean.split(' ', 1)[0]

    try:
        parts = clean.split(':')
        hour = int(parts[0])
        minute = int(parts[1]) if len(parts) > 1 else 0
        second = int(parts[2]) if len(parts) > 2 else 0

        if is_pm:
            if hour < 12:
                hour += 12
        elif is_am:
            if hour == 12:
                hour = 0

        return hour * 3600 + minute * 60 + second
    except (ValueError, IndexError):
        return None


def _seconds_to_display_time(seconds):
    """Convert seconds from midnight to 12-hour 'HH:MM AM/PM' string."""
    if seconds is None:
        return ''
    seconds = int(seconds) % 86400
    hours = seconds // 3600
    minutes = (seconds % 3600) // 60
    suffix = 'AM' if hours < 12 else 'PM'
    hour_12 = hours % 12 or 12
    return f"{hour_12:02d}:{minutes:02d} {suffix}"


def _split_course_code_name(course_name):
    """Split course string into code and title, e.g. 'IT101 - Intro to Computing' -> ('IT101', 'Intro to Computing')."""
    if not course_name:
        return 'TBA', ''
    c_str = str(course_name).strip()
    if ' - ' in c_str:
        parts = c_str.split(' - ', 1)
        return parts[0].strip(), parts[1].strip()
    match = re.match(r'^([A-Z]{2,5}[-\s]?\d{2,4}[A-Z]?)\s+(.*)$', c_str, re.IGNORECASE)
    if match:
        return match.group(1).strip(), match.group(2).strip()
    return c_str, ''


def generate_timetable_excel(schedule_type, entity_info, entries, timeslots=None, filter_metadata=None, theme='Blue'):
    """Generate a high-quality, vertically stretched Microsoft Excel (.xlsx) timetable workbook.

    Args:
        schedule_type: 'room', 'section', or 'professor'
        entity_info: dict or obj with name, type, department, etc.
        entries: list of schedule entry dictionaries
        timeslots: optional list of timeslot records from database
        filter_metadata: optional dict containing active filter parameters (semester, year, major, etc.)
        theme: optional theme name ('Blue', 'Red', 'Green', 'Purple', 'Orange', 'Teal', 'Pink')

    Returns:
        io.BytesIO: Binary Excel (.xlsx) stream
    """
    filter_metadata = filter_metadata or {}
    timeslots = timeslots or []
    entries = entries or []

    # Resolve Theme Palette
    theme_key = DEFAULT_THEME
    for t_name in EXCEL_THEMES:
        if t_name.lower() == str(theme).strip().lower():
            theme_key = t_name
            break
    palette = EXCEL_THEMES[theme_key]

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Class Schedule"

    # Define Theme Styles
    FONT_FAMILY = "Calibri"

    font_system = Font(name=FONT_FAMILY, size=9, bold=True, color=palette['accent_fg'])
    font_title = Font(name=FONT_FAMILY, size=15, bold=True, color=palette['header_bg'])
    font_subtitle = Font(name=FONT_FAMILY, size=11, bold=True, color=palette['accent_fg'])
    font_meta = Font(name=FONT_FAMILY, size=9, bold=False, color="334155")
    font_meta_bold = Font(name=FONT_FAMILY, size=9, bold=True, color=palette['header_bg'])

    font_header = Font(name=FONT_FAMILY, size=10, bold=True, color=palette['header_fg'])
    font_time_slot = Font(name=FONT_FAMILY, size=9, bold=True, color="1E293B")
    font_class_content = Font(name=FONT_FAMILY, size=9, bold=False, color="0F172A")
    font_lunch = Font(name=FONT_FAMILY, size=9, bold=True, color="94A3B8")

    fill_header = PatternFill(start_color=palette['header_bg'], end_color=palette['header_bg'], fill_type="solid")
    fill_time_col = PatternFill(start_color="F8FAFC", end_color="F8FAFC", fill_type="solid")
    fill_class = PatternFill(start_color=palette['class_fill'], end_color=palette['class_fill'], fill_type="solid")
    fill_lunch = PatternFill(start_color="F1F5F9", end_color="F1F5F9", fill_type="solid")
    fill_meta_bar = PatternFill(start_color=palette['meta_bar_fill'], end_color=palette['meta_bar_fill'], fill_type="solid")

    align_center = Alignment(horizontal="center", vertical="center", wrap_text=True)
    align_left_mid = Alignment(horizontal="left", vertical="center")
    align_right_mid = Alignment(horizontal="right", vertical="center")

    thin_border_color = palette['border_color']
    border_thin = Border(
        left=Side(style="thin", color=thin_border_color),
        right=Side(style="thin", color=thin_border_color),
        top=Side(style="thin", color=thin_border_color),
        bottom=Side(style="thin", color=thin_border_color),
    )
    border_header = Border(
        left=Side(style="thin", color=palette['header_border']),
        right=Side(style="thin", color=palette['header_border']),
        top=Side(style="medium", color=palette['header_border']),
        bottom=Side(style="medium", color=palette['header_border']),
    )

    # 1. Resolve Titles & Metadata
    school_year = filter_metadata.get('school_year') or "A.Y. 2026-2027"
    semester = (filter_metadata.get('semester') or '').strip()
    year_level = (str(filter_metadata.get('year') or '')).strip()
    major = (filter_metadata.get('major') or '').strip()
    program = (filter_metadata.get('program') or '').strip()

    if schedule_type == 'room':
        doc_heading = "ROOM SCHEDULE"
        room_name = entity_info.get('room_name') if isinstance(entity_info, dict) else str(entity_info)
        room_type = entity_info.get('room_type', 'Lecture') if isinstance(entity_info, dict) else 'Lecture'
        sub_heading = f"Room: {room_name}  ({room_type})"
        meta_parts = [
            f"Semester: {semester or 'All Semesters'}",
            f"Academic Year: {school_year}",
            f"Room Type: {room_type}",
            f"Program: {program or 'All Programs'}",
            f"Scheduled: {len(entries)} class{'es' if len(entries) != 1 else ''}",
        ]
    elif schedule_type == 'section':
        doc_heading = "SECTION SCHEDULE"
        sec_name = entity_info.get('section_name') or entity_info.get('section') if isinstance(entity_info, dict) else str(entity_info)
        sec_year = entity_info.get('year_level') if isinstance(entity_info, dict) else year_level
        sec_sem = entity_info.get('semester') if isinstance(entity_info, dict) else semester
        sec_maj = entity_info.get('major') if isinstance(entity_info, dict) else major
        sub_heading = f"Section: {sec_name}"
        meta_parts = [
            f"Year Level: {f'Year {sec_year}' if sec_year else (f'Year {year_level}' if year_level else 'N/A')}",
            f"Semester: {sec_sem or semester or 'All Semesters'}",
            f"Major: {sec_maj or major or 'General / Core'}",
            f"Academic Year: {school_year}",
            f"Weekly Classes: {len(entries)} class{'es' if len(entries) != 1 else ''}",
        ]
    elif schedule_type in ('teacher', 'professor'):
        doc_heading = "TEACHER SCHEDULE"
        if isinstance(entity_info, dict):
            p_first = entity_info.get('first_name', '')
            p_last = entity_info.get('last_name', '')
            prof_name = f"{p_first} {p_last}".strip() or entity_info.get('professor_name', 'Professor')
            dept = entity_info.get('department', 'N/A')
        else:
            prof_name = str(entity_info)
            dept = 'N/A'
        sub_heading = f"Professor: {prof_name}"
        meta_parts = [
            f"Department: {dept}",
            f"Semester: {semester or 'All Semesters'}",
            f"Academic Year: {school_year}",
            f"Program: {program or 'All Programs'}",
            f"Assigned Classes: {len(entries)} class{'es' if len(entries) != 1 else ''}",
        ]
    else:
        doc_heading = "CLASS SCHEDULE"
        sub_heading = "Schedule Details"
        meta_parts = [f"Semester: {semester or 'All Semesters'}", f"Academic Year: {school_year}"]

    # 2. Write Header (Rows 1 to 5)
    ws['A1'] = "CLASS SCHEDULING MANAGEMENT SYSTEM"
    ws['A1'].font = font_system

    ws['A2'] = doc_heading
    ws['A2'].font = font_title

    ws['A3'] = sub_heading
    ws['A3'].font = font_subtitle

    now_str = datetime.now().strftime("%b %d, %Y  %I:%M %p")
    ws['G1'] = f"Exported: {now_str}"
    ws['G1'].font = Font(name=FONT_FAMILY, size=8, color="64748B")
    ws['G1'].alignment = align_right_mid

    ws['G2'] = f"{palette['display_name']}  •  Official Timetable"
    ws['G2'].font = Font(name=FONT_FAMILY, size=9, bold=True, color=palette['accent_fg'])
    ws['G2'].alignment = align_right_mid

    # Metadata bar across Row 4
    ws.merge_cells('A4:G4')
    meta_bar_cell = ws['A4']
    meta_bar_cell.value = "   •   ".join(meta_parts)
    meta_bar_cell.font = font_meta_bold
    meta_bar_cell.fill = fill_meta_bar
    meta_bar_cell.alignment = align_center

    for col in range(1, 8):
        c_cell = ws.cell(row=4, column=col)
        c_cell.border = border_thin
        c_cell.fill = fill_meta_bar

    ws.row_dimensions[1].height = 18
    ws.row_dimensions[2].height = 24
    ws.row_dimensions[3].height = 20
    ws.row_dimensions[4].height = 22
    ws.row_dimensions[5].height = 8  # Spacer row

    # 3. Time Window & Intervals
    earliest_sec = 7 * 3600   # 7:00 AM
    latest_sec = 19 * 3600    # 7:00 PM
    lunch_sec = 12 * 3600     # 12:00 PM

    for ts in timeslots:
        s_val = _parse_time_to_seconds(ts.get('start_time'))
        e_val = _parse_time_to_seconds(ts.get('end_time'))
        l_val = _parse_time_to_seconds(ts.get('lunch_time'))
        if s_val is not None:
            earliest_sec = min(earliest_sec, s_val)
        if e_val is not None and e_val > earliest_sec:
            latest_sec = max(latest_sec, e_val)
        if l_val is not None:
            lunch_sec = l_val

    # Normalize parsed entries
    parsed_entries = []
    for e in entries:
        raw_start = e.get('start_time_raw') or e.get('class_start') or e.get('start') or e.get('start_time')
        raw_end = e.get('end_time_raw') or e.get('class_end') or e.get('end') or e.get('end_time')
        st_sec = _parse_time_to_seconds(raw_start)
        et_sec = _parse_time_to_seconds(raw_end)
        if st_sec is not None and et_sec is not None and et_sec > st_sec:
            earliest_sec = min(earliest_sec, st_sec)
            latest_sec = max(latest_sec, et_sec)
            c_code, c_title = _split_course_code_name(e.get('course_name'))
            parsed_entries.append({
                'day': (e.get('day') or '').strip().title(),
                'start_sec': st_sec,
                'end_sec': et_sec,
                'start_fmt': _seconds_to_display_time(st_sec),
                'end_fmt': _seconds_to_display_time(et_sec),
                'course_code': e.get('course_code') or c_code,
                'course_name': c_title or (c_code if not c_title else ''),
                'raw_course': e.get('course_name') or 'TBA',
                'professor': e.get('professor') or e.get('professor_name') or 'TBA',
                'room': e.get('room') or e.get('room_name') or 'TBA',
                'section': e.get('section') or 'TBA',
                'session_type': e.get('session_type') or 'Lecture',
            })

    # Hourly slot generation
    hourly_slots = []
    slot_curr = earliest_sec
    while slot_curr + 3600 <= latest_sec:
        slot_next = slot_curr + 3600
        is_lunch = (slot_curr <= lunch_sec < slot_next)
        s_lbl = f"{_seconds_to_display_time(slot_curr)} - {_seconds_to_display_time(slot_next)}"
        hourly_slots.append({
            'start_sec': slot_curr,
            'end_sec': slot_next,
            'label': s_lbl,
            'is_lunch': is_lunch,
        })
        slot_curr = slot_next

    if not hourly_slots:
        for h in range(7, 19):
            s_sec = h * 3600
            e_sec = (h + 1) * 3600
            hourly_slots.append({
                'start_sec': s_sec,
                'end_sec': e_sec,
                'label': f"{_seconds_to_display_time(s_sec)} - {_seconds_to_display_time(e_sec)}",
                'is_lunch': (h == 12),
            })

    days = ['Monday', 'Tuesday', 'Wednesday', 'Thursday', 'Friday', 'Saturday']

    # 4. Write Timetable Column Headers (Row 6)
    table_start_row = 6
    ws.row_dimensions[table_start_row].height = 28

    headers = ["TIME"] + [d.upper() for d in days]
    for col_idx, h_text in enumerate(headers, start=1):
        cell = ws.cell(row=table_start_row, column=col_idx)
        cell.value = h_text
        cell.font = font_header
        cell.fill = fill_header
        cell.alignment = align_center
        cell.border = border_header

    # Pre-map entries by day
    entries_by_day = {d: [] for d in days}
    for pe in parsed_entries:
        if pe['day'] in entries_by_day:
            entries_by_day[pe['day']].append(pe)

    num_slots = len(hourly_slots)
    grid_matches = [[[] for _ in range(len(days))] for _ in range(num_slots)]

    for d_idx, day in enumerate(days):
        for s_idx, slot in enumerate(hourly_slots):
            s_start = slot['start_sec']
            s_end = slot['end_sec']
            for pe in entries_by_day[day]:
                if pe['start_sec'] < s_end and s_start < pe['end_sec']:
                    grid_matches[s_idx][d_idx].append(pe)

    # 5. Spanning / Multi-hour Merges
    spanned_cells = set()
    merges = []

    for s_idx, slot in enumerate(hourly_slots):
        for d_idx, day in enumerate(days):
            col_idx = d_idx + 2  # Col 1 is Time, Col 2 is Monday...
            if (s_idx, d_idx) in spanned_cells:
                continue

            matches = grid_matches[s_idx][d_idx]
            if len(matches) == 1:
                single_entry = matches[0]
                end_s_idx = s_idx
                while end_s_idx + 1 < num_slots:
                    next_matches = grid_matches[end_s_idx + 1][d_idx]
                    if len(next_matches) == 1 and next_matches[0] == single_entry:
                        end_s_idx += 1
                    else:
                        break

                if end_s_idx > s_idx:
                    for r_sub in range(s_idx + 1, end_s_idx + 1):
                        spanned_cells.add((r_sub, d_idx))
                    start_row = table_start_row + 1 + s_idx
                    end_row = table_start_row + 1 + end_s_idx
                    merges.append((start_row, col_idx, end_row, col_idx))

    # 6. Populate Timetable Data Rows (Rows 7+) with VERTICAL STRETCHING
    ROW_HEIGHT_STRETCHED = 70  # Generous height for spacious vertical stretching

    for s_idx, slot in enumerate(hourly_slots):
        current_row = table_start_row + 1 + s_idx
        # Set large row height for vertical stretching
        ws.row_dimensions[current_row].height = ROW_HEIGHT_STRETCHED

        # Time Column (Col 1)
        time_cell = ws.cell(row=current_row, column=1)
        time_cell.value = slot['label']
        time_cell.font = font_time_slot
        time_cell.fill = fill_time_col
        time_cell.alignment = align_center
        time_cell.border = border_thin

        # Days Columns (Cols 2 to 7)
        for d_idx, day in enumerate(days):
            col_idx = d_idx + 2
            cell = ws.cell(row=current_row, column=col_idx)
            cell.border = border_thin

            if (s_idx, d_idx) in spanned_cells:
                # Subordinate cell in span: leave blank, but style border
                cell.fill = fill_class
                continue

            matches = grid_matches[s_idx][d_idx]
            if matches:
                cell.fill = fill_class
                cell.alignment = align_center

                card_texts = []
                for m in matches:
                    c_title_part = f" - {m['course_name']}" if m['course_name'] else ""
                    line1 = f"{m['course_code']}{c_title_part}"

                    if schedule_type == 'room':
                        line2 = f"Section: {m['section']}\nProf: {m['professor']}"
                    elif schedule_type == 'section':
                        line2 = f"Prof: {m['professor']}\nRoom: {m['room']}"
                    elif schedule_type in ('teacher', 'professor'):
                        line2 = f"Section: {m['section']}\nRoom: {m['room']}"
                    else:
                        line2 = f"Sec: {m['section']} | Room: {m['room']}"

                    line3 = f"[{m['session_type']}]  {m['start_fmt']} - {m['end_fmt']}"
                    card_texts.append(f"{line1}\n{line2}\n{line3}")

                cell.value = "\n--------------------\n".join(card_texts)
                cell.font = font_class_content

            elif slot['is_lunch']:
                cell.fill = fill_lunch
                cell.value = "LUNCH BREAK"
                cell.font = font_lunch
                cell.alignment = align_center
            else:
                cell.value = ""
                cell.alignment = align_center

    # Apply merges and fix merged cell borders
    for start_r, start_c, end_r, end_c in merges:
        ws.merge_cells(start_row=start_r, start_column=start_c, end_row=end_r, end_column=end_c)
        for r_m in range(start_r, end_r + 1):
            for c_m in range(start_c, end_c + 1):
                ws.cell(row=r_m, column=c_m).border = border_thin

    # 7. Column Widths
    ws.column_dimensions['A'].width = 20  # Time column
    for col_idx in range(2, 8):
        c_letter = get_column_letter(col_idx)
        ws.column_dimensions[c_letter].width = 26  # Days columns

    # 8. Page Setup & Printing Formatting
    ws.sheet_properties.tabColor = palette['header_bg']
    ws.page_setup.orientation = ws.ORIENTATION_LANDSCAPE
    ws.page_setup.paperSize = ws.PAPERSIZE_LETTER
    ws.sheet_properties.pageSetUpPr.fitToPage = True
    ws.page_setup.fitToWidth = 1
    ws.page_setup.fitToHeight = 0
    ws.views.sheetView[0].showGridLines = True

    # Freeze Header & Time Column (Panes frozen at B7)
    ws.freeze_panes = 'B7'

    buffer = io.BytesIO()
    wb.save(buffer)
    buffer.seek(0)
    return buffer
